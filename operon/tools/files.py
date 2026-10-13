"""Shared-drive capability: list, read (PDF, spreadsheets, text), write and move files.

Every path is resolved and confined to the shared-drive root; there is no delete.
"""

from __future__ import annotations

import csv
import io
import shutil
from pathlib import Path

from .base import Tool, ToolError, ToolResult, schema

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".log", ".tsv"}
WRITE_SUFFIXES = {".txt", ".md", ".csv", ".json"}
MAX_CHARS = 20000


class SharedDrive:
    def __init__(self, root: Path, read_only: bool = False) -> None:
        self.root = root.resolve()
        self.read_only = read_only

    def resolve(self, rel: str) -> Path:
        rel = (rel or ".").replace("\\", "/").strip()
        path = (self.root / rel.lstrip("/")).resolve()
        if path != self.root and not path.is_relative_to(self.root):
            raise ToolError(f"'{rel}' is outside the shared drive.")
        return path

    def rel(self, path: Path) -> str:
        return path.relative_to(self.root).as_posix() or "."

    def list(self, rel: str = ".") -> ToolResult:
        path = self.resolve(rel)
        if not path.exists():
            raise ToolError(f"'{rel}' does not exist.")
        if path.is_file():
            return ToolResult(f"{self.rel(path)} ({path.stat().st_size} bytes)")
        entries = []
        for child in sorted(path.iterdir()):
            if child.is_dir():
                count = sum(1 for _ in child.iterdir())
                entries.append(f"{self.rel(child)}/  ({count} items)")
            else:
                entries.append(f"{self.rel(child)}  ({child.stat().st_size} bytes)")
        return ToolResult(f"Contents of {self.rel(path)}/:\n" + ("\n".join(entries) or "(empty)"))

    def read(self, rel: str) -> ToolResult:
        path = self.resolve(rel)
        if not path.is_file():
            raise ToolError(f"'{rel}' is not a file. Use files_list to see what exists.")
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text = _read_pdf(path)
        elif suffix in (".xlsx", ".xlsm"):
            text = _read_xlsx(path)
        elif suffix in TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8", errors="replace")
        else:
            raise ToolError(f"Cannot read {suffix} files as text.")
        if len(text) > MAX_CHARS:
            text = text[:MAX_CHARS] + f"\n… (truncated {len(text) - MAX_CHARS} characters)"
        return ToolResult(f"File {self.rel(path)}:\n{text}")

    def write(self, rel: str, content: str, overwrite: bool = False) -> ToolResult:
        self._writable()
        path = self.resolve(rel)
        if path.suffix.lower() not in WRITE_SUFFIXES:
            raise ToolError(f"Only {sorted(WRITE_SUFFIXES)} files can be written.")
        if path.exists() and not overwrite:
            raise ToolError(f"{self.rel(path)} already exists; pass overwrite=true to replace it.")
        if path.suffix.lower() == ".csv":
            _check_csv(content)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8", newline="")
        return ToolResult(f"Wrote {self.rel(path)} ({len(content.encode())} bytes).", data={"file": self.rel(path)})

    def move(self, src: str, dst: str) -> ToolResult:
        self._writable()
        source, target = self.resolve(src), self.resolve(dst)
        if not source.is_file():
            raise ToolError(f"'{src}' is not a file.")
        if target.is_dir() or dst.endswith("/"):
            target.mkdir(parents=True, exist_ok=True)
            target = target / source.name
        if target.exists():
            raise ToolError(f"{self.rel(target)} already exists.")
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(source), str(target))
        return ToolResult(f"Moved {self.rel(source)} → {self.rel(target)}.", data={"file": self.rel(target)})

    def _writable(self) -> None:
        if self.read_only:
            raise ToolError("This session is read-only.")


def _read_pdf(path: Path) -> str:
    import pymupdf

    with pymupdf.open(path) as doc:
        pages = [f"--- page {i + 1} ---\n{page.get_text().strip()}" for i, page in enumerate(doc)]
    return "\n".join(pages)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"=== sheet: {ws.title} ===")
        for row in ws.iter_rows(values_only=True):
            if all(v is None for v in row):
                continue
            out.append("| " + " | ".join("" if v is None else str(v) for v in row) + " |")
    wb.close()
    return "\n".join(out)


def _check_csv(content: str) -> None:
    rows = list(csv.reader(io.StringIO(content)))
    if not rows:
        raise ToolError("CSV content is empty.")
    width = len(rows[0])
    bad = [i + 1 for i, r in enumerate(rows) if r and len(r) != width]
    if bad:
        raise ToolError(f"CSV rows {bad[:5]} have a different number of columns than the header ({width}). "
                        "Quote values that contain commas.")


def file_tools(drive: SharedDrive) -> list[Tool]:
    tools = [
        Tool("files_list", "List a folder on the shared drive (default: the root).",
             schema({"path": {"type": "string", "description": "Folder, e.g. ap_inbox"}}),
             lambda a: drive.list(a.get("path", "."))),
        Tool("files_read", "Read a file from the shared drive. PDFs and spreadsheets (.xlsx) are converted to text.",
             schema({"path": {"type": "string"}}, ["path"]),
             lambda a: drive.read(a["path"])),
    ]
    if not drive.read_only:
        tools += [
            Tool("files_write", "Create a text file (.csv, .md, .txt, .json) on the shared drive, e.g. a report.",
                 schema({"path": {"type": "string"}, "content": {"type": "string"},
                         "overwrite": {"type": "boolean"}}, ["path", "content"]),
                 lambda a: drive.write(a["path"], a["content"], bool(a.get("overwrite", False))),
                 idempotent=False),
            Tool("files_move", "Move a file to another folder on the shared drive (e.g. ap_inbox/x.pdf → ap_processed/).",
                 schema({"source": {"type": "string"}, "destination": {"type": "string",
                         "description": "Target folder (ending in /) or full target path"}},
                        ["source", "destination"]),
                 lambda a: drive.move(a["source"], a["destination"]), idempotent=False),
        ]
    return tools
